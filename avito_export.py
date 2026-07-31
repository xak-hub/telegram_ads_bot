import os
import shutil

from openpyxl import load_workbook

from avito_row import build_values

BASE_DIR = os.path.dirname(__file__)
TEMPLATE_PATH = os.path.join(BASE_DIR, "avito_template_base.xlsx")
EXPORT_PATH = os.path.join(BASE_DIR, "avito_export.xlsx")

SHEET_NAME = "Объявления"
FIRST_DATA_ROW = 5  # строки 1-4 — служебные, менять нельзя (см. лист "Инструкция")


def _ensure_export_file() -> None:
    if not os.path.exists(EXPORT_PATH):
        shutil.copy(TEMPLATE_PATH, EXPORT_PATH)


def _header_map(ws) -> dict:
    headers = {}
    for c in range(1, ws.max_column + 1):
        name = ws.cell(row=2, column=c).value
        if name:
            headers[name] = c
    return headers


def _first_empty_row(ws, id_col: int) -> int:
    r = FIRST_DATA_ROW
    while ws.cell(row=r, column=id_col).value:
        r += 1
    return r


def list_listing_ids() -> list[str]:
    """Возвращает ID всех объявлений, лежащих сейчас в фиде (в порядке строк)."""
    _ensure_export_file()
    wb = load_workbook(EXPORT_PATH)
    ws = wb[SHEET_NAME]
    id_col = _header_map(ws)["Уникальный идентификатор объявления"]
    ids = []
    r = FIRST_DATA_ROW
    while ws.cell(row=r, column=id_col).value:
        ids.append(str(ws.cell(row=r, column=id_col).value))
        r += 1
    return ids


def remove_listings(listing_ids) -> int:
    """Убирает из фида строки с указанными ID и возвращает число удалённых.

    Вместо ws.delete_rows (ломает валидаторы данных из шаблона Avito) переписываем
    выживших подряд с первой строки данных, а хвост чистим — так файл остаётся
    компактным (без «дыр», на которых обрывается _first_empty_row) и не портится."""
    ids = {str(x) for x in listing_ids}
    if not ids:
        return 0
    _ensure_export_file()
    wb = load_workbook(EXPORT_PATH)
    ws = wb[SHEET_NAME]
    headers = _header_map(ws)
    id_col = headers["Уникальный идентификатор объявления"]
    ncols = ws.max_column

    rows = []
    r = FIRST_DATA_ROW
    while ws.cell(row=r, column=id_col).value:
        rows.append([ws.cell(row=r, column=c).value for c in range(1, ncols + 1)])
        r += 1
    last_row = r - 1

    survivors = [vals for vals in rows if str(vals[id_col - 1]) not in ids]
    removed = len(rows) - len(survivors)
    if not removed:
        return 0

    # Присваиваем через .value: ws.cell(..., value=None) НЕ очищает ячейку —
    # None это значение по умолчанию параметра, и метод просто ничего не пишет.
    for rr in range(FIRST_DATA_ROW, last_row + 1):
        for c in range(1, ncols + 1):
            ws.cell(row=rr, column=c).value = None
    for i, vals in enumerate(survivors):
        rr = FIRST_DATA_ROW + i
        for c, val in enumerate(vals, start=1):
            ws.cell(row=rr, column=c).value = val

    wb.save(EXPORT_PATH)
    return removed


def append_listing(
    vision_params: dict, avito_answers: dict, price: str, address: str, listing_id: str,
    photo_urls: list[str] | None = None,
) -> None:
    """Добавляет объявление строкой в avito_export.xlsx — точную копию шаблона Avito."""
    _ensure_export_file()
    wb = load_workbook(EXPORT_PATH)
    ws = wb[SHEET_NAME]
    headers = _header_map(ws)

    row = _first_empty_row(ws, headers["Уникальный идентификатор объявления"])
    values = build_values(vision_params, avito_answers, price, address, listing_id, photo_urls)

    for col_name, value in values.items():
        col = headers.get(col_name)
        if col:
            ws.cell(row=row, column=col, value=value)

    wb.save(EXPORT_PATH)

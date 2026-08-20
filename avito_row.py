import datetime
import os
import re
from typing import Optional

from openpyxl import load_workbook

# Единый источник правды: список колонок читается напрямую из официального
# файла Avito, присланного пользователем ("Электроника - Ноутбуки - Шаблон
# 14-07-2026.xlsx"), а не хранится отдельным хардкодом — так Google Sheets
# и avito_export.xlsx гарантированно совпадают между собой и с реальной
# схемой Avito.

BASE_DIR = os.path.dirname(__file__)
TEMPLATE_PATH = os.path.join(BASE_DIR, "avito_template_base.xlsx")


def _load_columns() -> list[str]:
    wb = load_workbook(TEMPLATE_PATH, data_only=True)
    ws = wb["Объявления"]
    return [
        ws.cell(row=2, column=c).value
        for c in range(1, ws.max_column + 1)
        if ws.cell(row=2, column=c).value
    ]


ALL_COLUMNS = _load_columns()


def _slugify(text: str, max_len: int) -> str:
    if not text:
        return ""
    text = re.sub(r"[^0-9A-Za-zА-Яа-яЁё]+", "-", text).strip("-")
    return text[:max_len]


_STORAGE_STEPS = [128, 256, 512, 1024, 2048, 4096]


def round_storage(raw: str) -> str:
    """Округляет объём накопителя к ближайшему стандартному значению
    (128/256/512/1024/2048/4096): '494' -> '512'. Если распознать число не
    удалось — возвращает исходную строку как есть."""
    if not raw:
        return raw
    m = re.search(r"\d+", raw.replace(" ", ""))
    if not m:
        return raw
    val = int(m.group())
    nearest = min(_STORAGE_STEPS, key=lambda s: abs(s - val))
    return str(nearest)


def battery_health_percent(wear_raw: str) -> str:
    """Диагностические тулзы (AIDA64 и т.п.) показывают «степень изношенности» —
    для объявления понятнее «здоровье АКБ», это её зеркало: 88% износа -> 12%
    здоровья. Если число не распозналось — возвращает исходную строку как есть."""
    if not wear_raw:
        return wear_raw
    m = re.search(r"(\d+(?:[.,]\d+)?)", wear_raw)
    if not m:
        return wear_raw
    wear = float(m.group(1).replace(",", "."))
    health = max(0.0, 100.0 - wear)
    return f"{health:.0f}%"


def _short_cpu(cpu: str) -> str:
    if not cpu:
        return ""
    cpu = re.sub(r"\d+(\.\d+)?\s*GHz", "", cpu, flags=re.IGNORECASE)
    for word in ("Intel", "AMD", "Core", "Processor", "CPU", "(R)", "(TM)"):
        cpu = cpu.replace(word, "")
    return cpu.strip()


_CPU_TITLE_RE = re.compile(r"(i[3579])[\s-]?(\d{3,5}[a-z]{0,3})", re.IGNORECASE)


def _cpu_title_token(cpu: str) -> str:
    """Короткий вид процессора для заголовка, например "i5 8365u"."""
    if not cpu:
        return ""
    m = _CPU_TITLE_RE.search(cpu)
    if m:
        return f"{m.group(1).lower()} {m.group(2).lower()}"
    return _short_cpu(cpu)


# Линейка процессора для колонки Авито «Линейка процессора» — выводится из
# полного имени (например "Intel Core i7-1355U" -> "Core i7"). Значения
# соответствуют справочнику Авито для категории «Ноутбуки».
def cpu_line(cpu: str) -> str:
    """Выводит линейку процессора из полного имени:
    Core i3/i5/i7/i9, Ryzen 3/5/7/9, Celeron, Pentium, Xeon, Atom, Athlon,
    Apple M1/M2/M3 (+ Pro/Max/Ultra). Пустая строка, если не распознано."""
    if not cpu:
        return ""
    s = cpu.strip()

    m = re.search(r"Core\s+(i[3579])\b", s, re.IGNORECASE)
    if m:
        return f"Core {m.group(1).lower()}"

    # Новые линейки Intel: "Core Ultra 9 285H" -> "Core Ultra 9" (Ultra 5/7/9).
    m = re.search(r"Core\s+Ultra\s+([579])\b", s, re.IGNORECASE)
    if m:
        return f"Core Ultra {m.group(1)}"

    m = re.search(r"Ryzen\s+(\d)\b", s, re.IGNORECASE)
    if m:
        return f"Ryzen {m.group(1)}"

    m = re.search(r"\bM([1-4])(?:\s?(Pro|Max|Ultra))?\b", s, re.IGNORECASE)
    if m and ("apple" in s.lower() or "macbook" in s.lower() or s.lower().startswith("m")):
        suffix = (m.group(2) or "").title()
        return f"M{m.group(1)}{(' ' + suffix) if suffix else ''}"

    for line in ("Celeron", "Pentium", "Xeon", "Atom", "Athlon"):
        if re.search(rf"\b{line}\b", s, re.IGNORECASE):
            return line
    return ""


_DESC_INTRO_TAIL = (
    "Проверенный, настроенный и готовый к работе.\n"
    "С гарантией 6 месяцев и безлимитной постпродажной поддержкой.\n"
    "Установлен свежий Windows, последние драйвера, базовый софт.\n"
    "Проведен детейлинг внешки и узлов, обслужен.\n\n"
    "* Быстрая покупка самовыносом из офиса на Савёловской\n"
    "* Доставка по Москве в течении 3-х часов\n"
    "* Авито.Доставку любим, отправляем.\n"
    "* Любые формы оплаты"
)

_DESC_OUTRO = (
    "Зарядка, упаковочная коробка, товарный чек в комплекте.\n\n"
    "Всегда в наличии несколько единиц, можем укомплектовать ваш офис по индивидуальному запросу.\n\n"
    "Берём на абонентское обслуживание.\n"
    "На рынке электроники с 2005 года."
)


def _fmt_gb(value: str) -> str:
    """2048 -> '2 ТБ', 4096 -> '4 ТБ', 512 -> '512 ГБ', 48 -> '48 ГБ'."""
    try:
        v = float(str(value).strip())
    except (ValueError, TypeError):
        return str(value)
    if v >= 1024:
        tb = v / 1024
        tb_str = f"{tb:.0f}" if tb == int(tb) else f"{tb:.1f}"
        return f"{tb_str} ТБ"
    return f"{int(v)} ГБ"


def _gpu_vram_gb(vision_params: dict) -> str:
    """Валидный объём видеопамяти дискретной карты в ГБ ('2', '6') или ''.
    Принимает '2', '2 ГБ', '2GB'; отбрасывает мусор вне 1..32."""
    raw = str(vision_params.get("gpu_vram_gb", "")).strip()
    if not raw:
        return ""
    m = re.search(r"(\d+(?:[.,]\d+)?)", raw)
    if not m:
        return ""
    try:
        v = float(m.group(1).replace(",", "."))
    except ValueError:
        return ""
    if not (1 <= v <= 32):
        return ""
    return f"{int(v)}" if v == int(v) else f"{v:g}"


def _cpu_freq_suffix(vision_params: dict) -> str:
    """Вилка частот процессора: '(2.2-4.8 ГГц)'. Если известна только одна
    граница — '(до 4.8 ГГц)' / '(от 2.2 ГГц)'. Пусто — если частот нет."""
    ghz_min = str(vision_params.get("cpu_ghz_min", "")).strip().replace(",", ".").rstrip(".")
    ghz_max = str(vision_params.get("cpu_ghz_max", "")).strip().replace(",", ".").rstrip(".")
    # валидация: это должно быть число 0.3..6.0
    def _ok(v: str) -> bool:
        try:
            return 0.3 <= float(v) <= 6.0
        except ValueError:
            return False
    ghz_min = ghz_min if _ok(ghz_min) else ""
    ghz_max = ghz_max if _ok(ghz_max) else ""
    if ghz_min and ghz_max:
        if ghz_min == ghz_max:
            return f"({ghz_min} ГГц)"
        return f"({min(ghz_min, ghz_max, key=float)}-{max(ghz_min, ghz_max, key=float)} ГГц)"
    if ghz_max:
        return f"(до {ghz_max} ГГц)"
    if ghz_min:
        return f"(от {ghz_min} ГГц)"
    return ""


def _spec_block(vision_params: dict, avito_answers: dict) -> str:
    """Технический блок описания. Видеокарта и циклы АКБ — только если распознаны.
    Процессор — с вилкой частот (базовая..Turbo), память и накопитель —
    с максимумом апгрейда (SSD по умолчанию до 4 ТБ)."""
    cpu = vision_params.get("cpu", "").strip()
    gpu = vision_params.get("gpu", "").strip()
    ram = avito_answers.get("Объем оперативной памяти") or vision_params.get("ram_gb", "")
    storage = vision_params.get("storage_gb", "").strip()
    storage_type = avito_answers.get("Конфигурация накопителей", "").strip()
    screen = vision_params.get("screen_size", "").strip()
    resolution = vision_params.get("screen_resolution", "").strip()
    os_name = avito_answers.get("Операционная система") or vision_params.get("os", "")
    cycles = vision_params.get("battery_cycle_count", "").strip()

    # Максимумы апгрейда — только если больше текущего объёма (иначе бессмысленно).
    ram_max = vision_params.get("max_ram_gb", "").strip()
    try:
        if ram_max and ram and float(ram_max) <= float(str(ram).split()[0]):
            ram_max = ""
    except (ValueError, TypeError):
        ram_max = ""
    # SSD: если GLM не распознал максимум — по умолчанию 4 ТБ (современный M.2/NVMe).
    storage_max = vision_params.get("max_storage_gb", "").strip() or "4096"
    try:
        if storage and float(storage_max) <= float(round_storage(storage)):
            storage_max = ""
    except (ValueError, TypeError):
        storage_max = ""

    lines = []
    if cpu:
        cpu_line = f"Процессор: {cpu}"
        freq = _cpu_freq_suffix(vision_params)
        if freq:
            cpu_line += f" {freq}"
        lines.append(cpu_line)
    if ram:
        ram_line = f"Оперативная память: {ram} ГБ"
        if ram_max:
            ram_line += f" (расширение до {_fmt_gb(ram_max)})"
        lines.append(ram_line)
    if gpu:
        gpu_line = f"Видеокарта: {gpu}"
        vram = _gpu_vram_gb(vision_params)
        if vram:
            gpu_line += f" ({vram} ГБ)"
        lines.append(gpu_line)
    if storage:
        storage_line = f"Накопитель: {storage_type} {round_storage(storage)} ГБ".strip()
        if storage_max:
            storage_line += f" (расширение до {_fmt_gb(storage_max)})"
        lines.append(storage_line)
    if screen:
        screen_line = f"Экран: {screen}\""
        if resolution:
            screen_line += f", {resolution}px"
        lines.append(screen_line)
    if os_name:
        lines.append(f"ОС: {os_name}")
    if cycles:
        lines.append(f"Циклов АКБ: {cycles}")
    return "\n".join(lines)


def build_description(vision_params: dict, avito_answers: dict) -> str:
    """Полное маркетинговое описание объявления: вступление + техблок +
    (при наличии) заметка пользователя + заключение. Используется одинаково
    и для фида Avito, и для превью в Telegram."""
    brand = vision_params.get("brand", "").strip()
    model = vision_params.get("model", "").strip()
    title = f"{brand} {model}".strip() or "Ноутбук"

    intro = f"{title}.\n{_DESC_INTRO_TAIL}"
    parts = [intro, _spec_block(vision_params, avito_answers)]

    extra_note = vision_params.get("extra_note", "").strip()
    if extra_note:
        parts.append(extra_note)

    parts.append(_DESC_OUTRO)
    return "\n\n".join(p for p in parts if p)


def _build_title(vision_params: dict, avito_answers: dict) -> str:
    """Заголовок вида "Сенсорный Dell Rugged 7220 12" i5 8365u/8-ram/SSD256/LTE"."""
    brand = vision_params.get("brand", "").strip()
    model = vision_params.get("model", "").strip()
    screen = vision_params.get("screen_size", "").strip()
    cpu_token = _cpu_title_token(vision_params.get("cpu", ""))
    ram = avito_answers.get("Объем оперативной памяти") or vision_params.get("ram_gb", "")
    storage = vision_params.get("storage_gb", "").strip()
    storage_type = avito_answers.get("Конфигурация накопителей", "").strip()
    touchscreen = vision_params.get("touchscreen", "").strip().lower() == "да"
    lte = vision_params.get("lte", "").strip().lower() == "да"

    head_parts = []
    if touchscreen:
        head_parts.append("Сенсорный")
    brand_model = f"{brand} {model}".strip()
    if brand_model:
        head_parts.append(brand_model)
    if screen:
        head_parts.append(f'{screen}"')
    if cpu_token:
        head_parts.append(cpu_token)
    title = " ".join(head_parts) or "Ноутбук"

    tail_parts = []
    if ram:
        tail_parts.append(f"{ram}-ram")
    if storage:
        tail_parts.append(f"{storage_type}{round_storage(storage)}")
    if lte:
        tail_parts.append("LTE")
    if tail_parts:
        title += "/" + "/".join(tail_parts)

    return title[:100]


def make_listing_id(model: str = "", cpu: str = "", price: str = "") -> str:
    """Короткий читаемый ID вида Модель-Цена-Процессор (например T14s-6557-i5-1145G7).
    Временная метка в конце нужна для гарантии уникальности — модель и цена
    у двух разных объявлений вполне могут совпасть."""
    parts = [_slugify(model, 18), price, _slugify(_short_cpu(cpu), 14)]
    prefix = "-".join(p for p in parts if p) or "BOT"
    suffix = datetime.datetime.now().strftime("%H%M%S")
    return f"{prefix}-{suffix}"


def _truncate_words(text: str, limit: int) -> str:
    """Обрезает строку до limit символов по границе целого слова (не рубит
    посреди слова). Если первое же слово длиннее лимита — режет жёстко."""
    text = text.strip()
    if len(text) <= limit:
        return text
    cut = text[:limit]
    if " " in cut:
        cut = cut.rsplit(" ", 1)[0]
    return cut.rstrip(" ,")


def build_values(
    vision_params: dict,
    avito_answers: dict,
    price: str,
    address: str,
    listing_id: str,
    photo_urls: Optional[list[str]] = None,
) -> dict:
    """Собирает {название_колонки: значение} по реальным названиям колонок Avito."""
    brand = vision_params.get("brand", "")
    model = vision_params.get("model", "")
    # Заголовок = текст из «Вижу:» (проброшен из main.py), обрезанный по целым
    # словам под лимит Avito (50). Фолбэк — сгенерированный компактный заголовок.
    title_seen = vision_params.get("title_seen", "").strip()
    if title_seen:
        title = _truncate_words(title_seen, 50)
    else:
        title = _build_title(vision_params, avito_answers)[:50]

    description = build_description(vision_params, avito_answers)

    values = {
        "Уникальный идентификатор объявления": listing_id,
        "Адрес": address,
        "Название объявления": title,
        "Описание объявления": description[:3000],
        "Категория": "Ноутбуки",
        "Цена": price,
        "Производитель": brand,
        "Модель": model,
        "Линейка процессора": cpu_line(vision_params.get("cpu", "")),
        "Процессор": vision_params.get("cpu", ""),
        "Видеокарта": vision_params.get("gpu", ""),
        "Объем видеопамяти": _gpu_vram_gb(vision_params),
        "Диагональ экрана ноутбука": vision_params.get("screen_size", ""),
        "Общий объем накопителей": vision_params.get("storage_gb", ""),
        "Ссылки на фото": " | ".join(photo_urls) if photo_urls else "",
    }
    # Ответы из кнопок (Состояние, Вид объявления, аккумулятор, экран, корпус и т.д.)
    # уже используют точно такие же названия колонок — пишем как есть.
    values.update({k: v for k, v in avito_answers.items() if v})
    return values
